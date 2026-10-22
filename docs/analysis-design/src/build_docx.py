"""
Assemble the submission document from the rendered diagrams.

    python build_docx.py

Writes ../Project_Analysis_and_Design.docx. Wide diagrams go on landscape pages and
tall ones on portrait pages, each scaled to fill the page it is on.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
PNG = OUT / "png"
MARGIN = Cm(1.5)
INK = RGBColor(0x1F, 0x23, 0x28)
QUIET = RGBColor(0x57, 0x60, 0x6A)

FIGURES = []   # (number, caption) for the list of figures


def set_page(section, landscape):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    w, h = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
    section.page_width, section.page_height = w, h
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, MARGIN)


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Career Readiness Advisor · Project Analysis and Design · page ")
    run.font.size, run.font.color.rgb = Pt(9), QUIET
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


class Report:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(11)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 18), ("Heading 2", 13.5), ("Title", 26)):
            self.doc.styles[name].font.name = "Calibri"
            self.doc.styles[name].font.size = Pt(size)
            self.doc.styles[name].font.color.rgb = INK
        set_page(self.doc.sections[0], landscape=False)
        footer(self.doc.sections[0])
        self.landscape = False
        self.fig = 0

    def page(self, landscape):
        """Start a new page in the given orientation."""
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        set_page(s, landscape)
        self.landscape = landscape

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def para(self, text, size=None, color=None, italic=False, bold_lead=None):
        p = self.doc.add_paragraph()
        if bold_lead:
            r = p.add_run(bold_lead + " ")
            r.bold = True
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        r.italic = italic
        return p

    def bullets(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def figure(self, name, heading, text, caption, section=None):
        """A diagram on its own page, turned to suit its shape."""
        path = PNG / f"{name}.png"
        w, h = Image.open(path).size
        reserve = 3.6 + (3.6 if section else 0)     # headings, text and caption, in cm

        def fit(page_w, page_h):
            return min((page_w - 2 * 1.5) / w, (page_h - 2 * 1.5 - reserve) / h)

        # Whichever way round prints the diagram larger.
        landscape = fit(29.7, 21.0) > fit(21.0, 29.7)
        scale = fit(29.7, 21.0) if landscape else fit(21.0, 29.7)
        self.page(landscape)
        if section:
            self.h1(section[0])
            self.para(section[1])
        self.h2(heading)
        self.para(text)
        pic = self.doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(path), width=Cm(w * scale))
        self.fig += 1
        FIGURES.append((self.fig, caption))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.fig}. {caption}")
        r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), QUIET

    def table(self, header, rows, widths):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.bold, r.font.size = True, Pt(10)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:fill"), "E7EEF9")
            cell._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(v)
                r.font.size = Pt(10)
                if i == 0 and len(header) == 2:
                    r.bold = True
        for row in t.rows:
            for i, wd in enumerate(widths):
                row.cells[i].width = Cm(wd)
        self.doc.add_paragraph()

    def save(self, path):
        self.doc.save(path)


def use_case(rep, title, fields):
    rep.para("", size=4)
    p = rep.doc.add_paragraph()
    r = p.add_run(title)
    r.bold, r.font.size = True, Pt(12)
    rep.table(["Field", "Description"], fields, [3.6, 13.8])


def build():
    rep = Report()
    d = rep.doc

    # ---- cover ------------------------------------------------------------
    d.add_paragraph().paragraph_format.space_after = Pt(60)
    t = d.add_paragraph(style="Title")
    t.add_run("Career Readiness Advisor")
    rep.para("Project Analysis and Design", size=18, color=QUIET)
    rep.para("UML model · use case · activity · sequence · database · class · ER diagrams",
             size=12, color=QUIET)
    rep.para("17 October 2026", size=12, color=QUIET)
    d.add_paragraph().paragraph_format.space_after = Pt(24)
    rep.para("The Career Readiness Advisor is an iOS-styled Flutter app over a stateless FastAPI "
             "service. It compares a student's completed courses, and any skills they report "
             "from outside class, against what 82 Saudi job postings across eight Information "
             "Systems careers ask for. It returns a demand-weighted readiness figure, the "
             "missing skills ranked by demand, and the university and online courses that "
             "close them. Department staff see the same catalog read against the market.")
    rep.para("Every diagram below is drawn from the code and data as they stand on 17 October "
             "2026: names, fields, endpoints and counts are the real ones.")
    rep.h2("Contents")
    rep.bullets([
        "1. UML model: component, deployment and package diagrams",
        "2. Use case diagrams and use case descriptions",
        "3. Activity diagrams",
        "4. Sequence diagrams",
        "5. Database diagram",
        "6. Class diagrams",
        "7. ER diagram",
    ])

    # ---- 1 ----------------------------------------------------------------
    rep.figure("1a_component", "1a. Component diagram",
               "The app is layered: screens use Riverpod state, state uses repositories, and "
               "repositories call the API over HTTP. The API builds the engine state once at "
               "start-up from six data files, and only loads the embedding model when its "
               "cache misses.",
               "Component diagram of the app, the API, the engine and the data files",
               section=("1. UML model",
                        "The system has three parts: the iPhone app, a FastAPI service, and the "
                        "skill-gap engine it wraps. The app never touches the engine or its data "
                        "files directly: everything goes through seven REST endpoints, and the "
                        "two that read a student's profile store nothing. The three views show "
                        "the parts and their dependencies (component), where each part runs "
                        "(deployment), and how the app's code is organised (package)."))
    rep.figure("1b_deployment", "1b. Deployment diagram",
               "The app runs on an iPhone or the iOS Simulator and keeps the profile in the "
               "phone's own storage. The API runs as one Python process on the development "
               "machine, on port 8000. The outside services are used only at build time, at "
               "collection time, or when a student opens an online course.",
               "Deployment diagram of the phone, the server and the outside services")
    rep.figure("1c_packages", "1c. Package diagram",
               "The app's code is split into eight feature packages over a shared core. Each "
               "feature keeps only the layers it needs: home and settings are screens only, "
               "and planning reuses the gap call instead of having its own data layer.",
               "Package diagram of the app's feature packages")

    # ---- 2 ----------------------------------------------------------------
    rep.figure("2a_usecase_app", "2a. App use cases",
               "Creating and editing a profile share one form. Comparing careers only happens "
               "when a profile names two or three careers, and opening an online course hands "
               "the student to the browser.",
               "Use case diagram for students and department staff",
               section=("2. Use case diagrams",
                        "Two kinds of user work in the app: students and department staff. Both "
                        "are app users, so both can switch language. A third actor, the "
                        "researcher, works only with the offline tools that prepare the data the "
                        "app reads."))
    rep.figure("2b_usecase_data", "2b. Data-preparation use cases",
               "The researcher collects and labels the data, measures the matcher, and starts "
               "the API. The embedding model is loaded by the validation sweep, by the "
               "course-reach report's --explain option, and when the engine's cache misses.",
               "Use case diagram for the researcher's data-preparation tools")
    rep.page(False)
    rep.h2("2c. Use case descriptions")
    rep.para("The five use cases that carry the app's purpose, written out.")
    use_case(rep, "UC-1 Create profile", [
        ("Actor", "Student"),
        ("Precondition", "No profile is saved on this device"),
        ("Main flow", "1) The app opens on “Your profile” and loads the careers and courses from "
                      "the API. 2) The student picks 1 to 3 target careers and may type a major. "
                      "3) The student ticks completed courses, with search, “select all shown”, "
                      "a level shortcut, or review and clear. 4) “Next: other skills” opens the "
                      "skills step, where the student may tick skills learned outside class; "
                      "skills the chosen courses teach are marked and cannot be ticked. 5) “See "
                      "my readiness” saves the profile on the device and opens the home "
                      "dashboard."),
        ("Alternate flows", "1a) The API is unreachable: an error with “Try again”. 2a) A fourth "
                            "career is refused. 5a) The save fails: the skills step says so and "
                            "keeps the draft."),
        ("Postcondition", "The profile is saved on the device, and onboarding does not run again "
                          "while it exists"),
    ])
    use_case(rep, "UC-2 View full readiness", [
        ("Actor", "Student"),
        ("Precondition", "A profile is saved"),
        ("Main flow", "1) From home, the student taps “Full readiness”. 2) The app sends the "
                      "profile, read against the main career, to POST /gap. 3) The screen shows "
                      "the quality flag, the readiness percentage, covered skills split into "
                      "“from courses” and “self-reported”, and missing skills ranked by their "
                      "share of postings. 4) Where self-reported skills change the figure, the "
                      "courses-only figure stands beside it."),
        ("Extension", "With 2 or 3 careers, they stand side by side with score, covered and "
                      "missing counts, and posting count. The closest is marked when the scores "
                      "differ, and tapping one switches the details to it."),
        ("Alternate flows", "2a) The request fails: an error with “Try again”. 2b) The profile "
                            "names a course, skill or career the API no longer serves: it is "
                            "left out, and a dropped main career gives way to the next."),
        ("Postcondition", "Nothing is stored; the reading is computed for each request"),
    ])
    use_case(rep, "UC-3 Plan courses (what-if)", [
        ("Actor", "Student"),
        ("Precondition", "A profile is saved"),
        ("Main flow", "1) The student opens “Plan courses” from home or readiness. 2) The app "
                      "lists the courses not yet taken, by track, each with the current gaps it "
                      "would close. 3) The student ticks a course, and the plan is saved on the "
                      "device at once. 4) The app calls POST /gap again with completed and "
                      "planned courses together. 5) The projection stands beside today's "
                      "figure, with the points gained and the gaps the plan closes, each counted "
                      "once."),
        ("Alternate flows", "1a) With several careers, the student switches the career the plan "
                            "is read against. 3a) The save fails: the plan is put back and a "
                            "message says so."),
        ("Postcondition", "The planned courses are saved with the profile; a course later marked "
                          "completed leaves the plan"),
    ])
    use_case(rep, "UC-4 View suggested courses", [
        ("Actor", "Student; the web browser for the extension"),
        ("Precondition", "A profile is saved"),
        ("Main flow", "1) From readiness, the student taps “Suggested courses”. 2) The app sends "
                      "the profile to POST /recommendations. 3) University courses come first, "
                      "ranked by the demand they close, each marked as counting toward the "
                      "degree. 4) Online courses follow, only for skills no catalog course "
                      "teaches, each marked as not counting. 5) Missing skills with neither are "
                      "listed, not dropped."),
        ("Extension", "The student taps an online course and the browser opens it; if nothing "
                      "can open it, a message says so"),
        ("Postcondition", "Nothing is stored"),
    ])
    use_case(rep, "UC-5 View curriculum gaps", [
        ("Actor", "Department staff"),
        ("Precondition", "None; the department tab works with or without a student profile"),
        ("Main flow", "1) Staff open the “Department” tab, which is built the first time it "
                      "opens. 2) The sample-size caveat comes first, with the real posting "
                      "counts, then the quality flag. 3) GET /curriculum/gap returns coverage "
                      "(taught skills over demanded skills) and the threshold in postings. 4) The "
                      "screen lists demanded skills no course teaches, by demand, with the "
                      "careers asking; skills named by too few postings to count; taught skills "
                      "no posting names; and coverage for each career."),
        ("Extension", "Staff pick one career's pill to see that career alone"),
        ("Postcondition", "Nothing is stored"),
    ])

    # ---- 3 ----------------------------------------------------------------
    rep.figure("3a_activity_journey", "3a. The student's journey",
               "The app decides at launch: no saved profile means onboarding, and a saved one "
               "goes straight to home. Every reading after that is a round trip to POST /gap, "
               "and every option from home leads back to it.",
               "Activity diagram of the student's journey through the app",
               section=("3. Activity diagrams",
                        "Four activities carry the system: the student's journey through the "
                        "app, the gap calculation behind every readiness figure, the engine's "
                        "start-up, and what-if planning."))
    rep.figure("3b_activity_gap", "3b. The gap calculation (POST /gap)",
               "The request is checked before anything is computed, and an unknown career, "
               "course or skill ends it with HTTP 422. The gap is then read twice, "
               "independently, once with self-reported skills and once from courses alone, so "
               "the academic record is never blurred with the student's own claim.",
               "Activity diagram of the POST /gap calculation")
    rep.figure("3c_activity_startup", "3c. Engine start-up",
               "The server refuses to start on a stale config or a bad label rather than serve "
               "wrong figures. When the cache key matches, start-up skips the model entirely; "
               "only a changed input brings back the embedding work.",
               "Activity diagram of the engine start-up and profile building")
    rep.figure("3d_activity_planning", "3d. What-if planning",
               "A tick shows on screen before the save completes, and a failed save puts the "
               "plan back. The projection is the same gap calculation with planned courses "
               "counted as taken.",
               "Activity diagram of what-if course planning")

    # ---- 4 ----------------------------------------------------------------
    rep.figure("4a_sequence_onboarding", "4a. Onboarding and saving the profile",
               "The form's lists come from three GET requests. Finishing writes one JSON "
               "document to the phone's shared_preferences; nothing goes to the server.",
               "Sequence diagram of onboarding and saving the profile",
               section=("4. Sequence diagrams",
                        "Six interactions, traced through the real classes and endpoints. "
                        "Filled arrowheads are calls, dashed arrows are replies, and the bars "
                        "show which object is working."))
    rep.figure("4b_sequence_readiness", "4b. Reading readiness on home",
               "The dashboard first drops any id the API no longer serves, then posts the "
               "profile to /gap. The engine reads the gap twice, with self-reported skills and "
               "with courses alone.",
               "Sequence diagram of reading readiness on the home dashboard")
    rep.figure("4c_sequence_suggestions", "4c. Suggested courses",
               "University courses that close a gap come first and count toward the degree. "
               "Online courses are offered only for gaps no catalog course teaches, and open in "
               "the phone's browser.",
               "Sequence diagram of the suggested courses")
    rep.figure("4d_sequence_planning", "4d. Planning a course",
               "The plan changes on screen before the save finishes, so a second quick tap "
               "builds on the first; a failed save rolls it back. The projection is a second "
               "POST /gap with planned courses counted as taken.",
               "Sequence diagram of planning a course")
    rep.figure("4e_sequence_department", "4e. The department view",
               "The tab is built the first time it opens. Every career's view comes first, and "
               "a career pill repeats the same request for one career.",
               "Sequence diagram of the department's curriculum view")
    rep.figure("4f_sequence_startup", "4f. Server start-up",
               "The engine is built once per process. When the cache file matches the key, the "
               "model is never created; on a miss it is downloaded once, and the results are "
               "written back to the cache.",
               "Sequence diagram of the server start-up")

    # ---- 5 ----------------------------------------------------------------
    rep.page(False)
    rep.h1("5. Database diagram")
    rep.para("The system has no database server. The data is nine stores: eight CSV and JSON "
             "files on the server, six of which the engine reads once at start-up, and a "
             "key-value store on the phone for the student's profile. The server's SQLite "
             "profile database was removed on 1 October 2026, so no student data is kept on "
             "the server; this is deliberate, and accounts are future work.")
    rep.table(["Store", "Where", "Rows", "Written by", "Read by"], [
        ("taxonomy.json", "server, skillgap/", "91 skills", "hand", "engine, validate.py"),
        ("config.json", "server, skillgap/", "1 document", "validate.py", "engine start-up"),
        ("labeled_postings.csv", "server, test/", "82 postings, 15 labelled",
         "collect_postings.py, hand labels", "engine, validate.py"),
        ("labeled_postings_excluded.csv", "server, test/", "28 postings", "collect_postings.py",
         "the researcher"),
        ("labeled_courses.csv", "server, skillgap/", "33 courses", "hand", "engine, validate.py"),
        ("labeled_courses_excluded.csv", "server, skillgap/", "15 courses", "hand",
         "the researcher"),
        ("online_courses.csv", "server, skillgap/", "47 rows, 28 courses",
         "hand; online.py sync adds placeholders", "engine"),
        ("engine-<key>.json", "server, backend/.cache/", "8 profiles, 33 courses",
         "engine, on a cache miss", "engine start-up"),
        ("shared_preferences", "the iPhone", "2 keys", "the app", "the app"),
    ], [4.6, 3.4, 3.2, 3.6, 2.8])
    rep.para("The skill lists inside the CSV files are stored as one field of ids separated by "
             "“;”. Normalised into tables, each list would become a link table between its "
             "store and the taxonomy; the ER diagram in section 7 shows those relationships.")
    rep.figure("5_database_stores", "5a. The data stores",
               "Each box is a store and its fields. PK identifies a row, FK points at another "
               "store, and the crow's feet give each link's cardinality.",
               "Database diagram of the data files and the phone's key-value store")

    # ---- 6 ----------------------------------------------------------------
    rep.figure("6a_class_engine", "6a. The skill-gap engine (skillgap.py)",
               "The Taxonomy embeds every skill's English and Arabic labels with one Embedder. "
               "The module's functions turn postings into a CareerProfile and a student's "
               "skills into a GapResult.",
               "Class diagram of the skill-gap engine",
               section=("6. Class diagrams",
                        "Five class diagrams cover the code, from the engine to the screens. "
                        "Names, fields and signatures are taken from the source; + is public, - "
                        "is private, a filled diamond is composition, a hollow triangle is "
                        "generalisation, and a dashed arrow is a dependency."))
    rep.figure("6b_class_service", "6b. The service (engine.py, online.py, main.py)",
               "One EngineState holds the 33 courses, 8 career profiles and the online catalog. "
               "Every response is built from it, as a SourcedGap, Suggestions or "
               "CurriculumGap.",
               "Class diagram of the service's EngineState and its parts")
    rep.figure("6c_class_api", "6c. The API's request and response objects (schemas.py)",
               "StudentIn is the only request body. Each response that reports a figure "
               "carries ValidationInfo, and a recommendation is one of two shapes told apart by "
               "its source field.",
               "Class diagram of the API's request and response objects")
    rep.figure("6d_class_app_domain", "6d. The app's domain model",
               "The app's models mirror the API's responses in camelCase, keeping the fields "
               "the screens use. StudentProfile is the only one the app writes; ServedIds trims "
               "a saved profile to what the API still serves.",
               "Class diagram of the app's domain model")
    rep.figure("6e_class_app_layers", "6e. The app's layers",
               "Screens watch Riverpod notifiers and providers; providers call repositories, "
               "and repositories call one ApiClient. The profile goes through the ProfileStore "
               "interface, which the phone's shared_preferences implements and the tests "
               "replace with a fake.",
               "Class diagram of the app's screens, state and data access")

    # ---- 7 ----------------------------------------------------------------
    rep.figure("7_er_chen", "7a. Entities and relationships",
               "Rectangles are entities, diamonds relationships and ovals attributes; an "
               "underlined attribute is a key, a dashed one is derived, and a double line means "
               "every instance takes part.",
               "ER diagram in Chen notation",
               section=("7. ER diagram",
                        "SKILL is the hub of the model: postings, careers, courses, online "
                        "courses and the student all connect through it. There are seven "
                        "entities and ten relationships, all many-to-many except a posting's "
                        "career and a skill's category."))
    rep.page(False)
    rep.h2("7b. Relationships")
    rep.table(["Relationship", "Between", "Cardinality", "Meaning"], [
        ("BELONGS_TO", "JOB_POSTING, CAREER", "N : 1", "Every posting is collected for one career"),
        ("DEMANDS", "CAREER, SKILL", "M : N", "A career's demand for a skill: the share of its "
                                              "postings naming it (derived)"),
        ("LABELLED_WITH", "JOB_POSTING, SKILL", "M : N", "Hand labels, on the 15 evaluation "
                                                         "postings only"),
        ("IN_CATEGORY", "SKILL, SKILL_CATEGORY", "N : 1", "Each of the 91 skills is in one of 8 "
                                                          "categories"),
        ("TEACHES", "COURSE, SKILL", "M : N", "The served course mapping; mapped_by is labels or "
                                              "extraction"),
        ("COVERS", "ONLINE_COURSE, SKILL", "M : N", "An online course offered for a skill no "
                                                    "catalog course teaches"),
        ("TARGETS", "STUDENT_PROFILE, CAREER", "M : N", "1 to 3 careers per student; rank 1 is "
                                                        "the main career"),
        ("COMPLETED", "STUDENT_PROFILE, COURSE", "M : N", "Courses the student has taken"),
        ("PLANS", "STUDENT_PROFILE, COURSE", "M : N", "Courses in the what-if plan"),
        ("REPORTS", "STUDENT_PROFILE, SKILL", "M : N", "Skills the student holds from outside "
                                                       "coursework; unverified"),
    ], [3.2, 4.6, 2.4, 7.2])

    rep.h2("List of figures")
    for n, cap in FIGURES:
        rep.para(f"Figure {n}. {cap}", size=10)
    rep.save(OUT / "Project_Analysis_and_Design.docx")
    print(f"{len(FIGURES)} figures")


if __name__ == "__main__":
    build()
